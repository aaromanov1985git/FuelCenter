"""
Тесты API резервуаров, лимитов карт и отчёта «Заправки по картам»
"""
import json
from datetime import datetime, timedelta
from decimal import Decimal

import pytest

from app.models import CardLimit, Provider, ProviderTemplate, Tank, TankReading, TopazSyncState, Transaction

SOURCE_NOW = datetime(2026, 9, 14, 14, 0, 0)


@pytest.fixture
def mazs(test_db):
    provider = Provider(name="МАЗС", code="MAZS")
    test_db.add(provider)
    test_db.flush()
    template = ProviderTemplate(
        provider_id=provider.id,
        name="MAZS",
        connection_type="firebird",
        field_mapping="{}",
        connection_settings=json.dumps({"host": "10.30.1.8", "database": "db.fdb"}),
        fuel_type_mapping=json.dumps({"Бензин": "АИ-92", "ДТ3": "ДТ"}),
    )
    test_db.add(template)
    test_db.flush()
    # Сервер Топаза живёт на 5 часов впереди сервера GSM
    run_at = datetime.now()
    test_db.add(TopazSyncState(
        template_id=template.id, source_kind="snapshots", last_tank_row_id=0,
        last_run_at=run_at, source_clock=run_at + timedelta(hours=5), last_status="success",
    ))
    source_now = run_at + timedelta(hours=5)

    def tank(key, num, fuel, volume, density, measured_at, capacity=None, **extra):
        item = Tank(
            provider_id=provider.id, template_id=template.id, azs_code="1016201", source_key=key,
            tank_number=num, source_name=f"Емкость {num} - 1016201", source_fuel=fuel,
            last_volume=Decimal(str(volume)), last_density=Decimal(str(density)), last_mass=Decimal("1000"),
            last_measured_at=measured_at, capacity_liters=Decimal(str(capacity)) if capacity else None, **extra,
        )
        test_db.add(item)
        return item

    petrol = tank("snap:1", 1, "АИ-92", 5409.18, 765.06, source_now - timedelta(minutes=10), capacity=30000)
    diesel = tank("snap:3", 2, "ДТ", 4230.71, 826.81, source_now - timedelta(minutes=5), capacity=30000)
    mirror = tank("snap:2", 1, "ДТ", 5477.05, 765.05, source_now - timedelta(days=3), capacity=30000)
    mirror.azs_code = "807211"
    test_db.commit()
    return {"provider": provider, "template": template, "petrol": petrol, "diesel": diesel, "mirror": mirror}


class TestTankOverview:
    def test_groups_by_station_and_sums_fuel(self, client, auth_headers, test_db, mazs):
        from app.models import GasStation
        station = GasStation(provider_id=mazs["provider"].id, original_name="1016201", name="1016201", azs_number="1016201",
                             location='База АО "УТТ"', settlement="Нягань", region="ХМАО-Ю")
        test_db.add(station)
        test_db.flush()
        for tank in (mazs["petrol"], mazs["diesel"]):
            tank.gas_station_id = station.id
        test_db.commit()

        response = client.get("/api/v1/tanks", headers=auth_headers)
        assert response.status_code == 200
        body = response.json()
        assert body["total_tanks"] == 3
        stations = {s["azs_code"]: s for s in body["stations"]}
        main = stations["1016201"]
        assert (main["location"], main["settlement"], main["region"]) == ('База АО "УТТ"', "Нягань", "ХМАО-Ю")
        fuels = {f["fuel_type"]: f for f in main["fuels"]}
        assert fuels["АИ-92"]["volume"] == 5409.18
        assert fuels["АИ-92"]["fill_percent"] == 18.0
        assert fuels["ДТ"]["capacity_liters"] == 30000
        assert body["sync"][0]["template_name"] == "MAZS"

    def test_age_uses_source_clock_and_flags_problems(self, client, auth_headers, mazs):
        body = client.get("/api/v1/tanks", headers=auth_headers).json()
        tanks = {t["source_key"]: t for s in body["stations"] for t in s["tanks"]}
        assert tanks["snap:1"]["age_minutes"] < 30
        assert tanks["snap:1"]["warnings"] == []
        # «ДТ» с бензиновой плотностью и замером трёхдневной давности
        assert set(tanks["snap:2"]["warnings"]) == {"density_mismatch", "stale"}

    def test_fuel_age_is_freshest_with_oldest_alongside(self, client, auth_headers, test_db, mazs):
        source_now = datetime.now() + timedelta(hours=5)
        second = Tank(provider_id=mazs["provider"].id, template_id=mazs["template"].id, azs_code="1016201",
                      source_key="snap:9", tank_number=9, source_fuel="АИ-92", last_volume=Decimal("1000"),
                      last_density=Decimal("765"), last_measured_at=source_now - timedelta(hours=17))
        test_db.add(second)
        test_db.commit()

        body = client.get("/api/v1/tanks", headers=auth_headers).json()
        station = next(s for s in body["stations"] if s["azs_code"] == "1016201")
        petrol = next(f for f in station["fuels"] if f["fuel_type"] == "АИ-92")
        assert petrol["age_minutes"] < 30
        assert 17 * 60 - 5 <= petrol["oldest_age_minutes"] <= 17 * 60 + 5

    def test_inactive_tank_excluded_from_station_totals(self, client, auth_headers, admin_auth_headers, mazs):
        mirror_id = mazs["mirror"].id
        response = client.patch(f"/api/v1/tanks/{mirror_id}", json={"is_active": False}, headers=admin_auth_headers)
        assert response.status_code == 200
        assert response.json()["is_active"] is False

        body = client.get("/api/v1/tanks", headers=auth_headers).json()
        station = next(s for s in body["stations"] if s["azs_code"] == "807211")
        assert station["fuels"] == []


class TestSessionEstimate:
    def test_overflow_pair_uses_shift_reading_minus_dispensed(self, client, auth_headers, test_db):
        from app.models import UploadEvent
        from app.services.topaz_sync_service import EVENT_ROW_OFFSET

        provider = Provider(name="КАЗС", code="KAZS")
        test_db.add(provider)
        test_db.flush()
        template = ProviderTemplate(provider_id=provider.id, name="KAZS", connection_type="firebird", field_mapping="{}",
                                    connection_settings="{}", fuel_type_mapping=json.dumps({"ДТ3": "ДТ", "Бензин": "АИ-92"}))
        test_db.add(template)
        test_db.flush()
        run_at = datetime.now()
        test_db.add(TopazSyncState(template_id=template.id, source_kind="sessions", last_tank_row_id=0,
                                   last_run_at=run_at, source_clock=run_at + timedelta(hours=5), last_status="success"))
        source_now = run_at + timedelta(hours=5)
        shift = source_now.replace(hour=0, minute=0, second=0, microsecond=0)
        dispense_at = shift + timedelta(hours=14, minutes=36)
        if dispense_at > source_now:
            dispense_at = source_now - timedelta(minutes=5)

        def tank(num, volume_at_shift, last_volume, last_at):
            item = Tank(provider_id=provider.id, template_id=template.id, azs_code="505221", source_key=f"ses:505221:{num}",
                        tank_number=num, source_fuel="ДТ1", capacity_liters=Decimal("10000"),
                        last_volume=Decimal(str(last_volume)), last_density=Decimal("834"), last_measured_at=last_at)
            test_db.add(item)
            test_db.flush()
            test_db.add(TankReading(tank_id=item.id, measured_at=shift, volume=Decimal(str(volume_at_shift)), source_row_id=11400 + num))
            return item

        first = tank(1, 7842.16, 7574.46, dispense_at)  # перелив из рез. 2 уже внутри этого замера
        tank(2, 7835.08, 7835.08, shift)
        tank(3, 7028.69, 7028.69, shift)
        test_db.add(TankReading(tank_id=first.id, measured_at=dispense_at, volume=Decimal("7574.46"),
                                source_row_id=EVENT_ROW_OFFSET + 112835))
        for qty in ("113", "120.05", "30.09", "250.07", "40.08"):
            test_db.add(Transaction(provider_id=provider.id, azs_number="505221", product="ДТ", card_number="К051",
                                    transaction_date=dispense_at - timedelta(minutes=1), quantity=Decimal(qty)))
        test_db.add(Transaction(provider_id=provider.id, azs_number="505221", product="ДТ", card_number="К051",
                                transaction_date=shift - timedelta(hours=1), quantity=Decimal("500")))  # до смены
        test_db.add(UploadEvent(source_type="auto", status="success", provider_id=provider.id, template_id=template.id,
                                created_at=run_at - timedelta(minutes=20)))
        test_db.commit()

        body = client.get("/api/v1/tanks", params={"provider_id": provider.id}, headers=auth_headers).json()
        diesel = body["stations"][0]["fuels"][0]
        assert diesel["fuel_type"] == "ДТ"
        assert diesel["estimate_base_volume"] == 22705.93
        assert diesel["estimate_dispensed"] == 553.29
        assert diesel["volume"] == 22152.64
        assert diesel["fill_percent"] == 73.8
        assert 15 <= diesel["age_minutes"] <= 25
        assert diesel["oldest_age_minutes"] is None


class TestTankUpdate:
    def test_admin_sets_capacity_override_and_group(self, client, admin_auth_headers, mazs):
        tank_id = mazs["diesel"].id
        response = client.patch(
            f"/api/v1/tanks/{tank_id}",
            json={"capacity_liters": 10000, "fuel_type_override": "ДТ", "overflow_group": "A"},
            headers=admin_auth_headers,
        )
        assert response.status_code == 200
        body = response.json()
        assert body["capacity_liters"] == 10000
        assert body["fuel_type"] == "ДТ"
        assert body["overflow_group"] == "A"
        assert body["fill_percent"] == 42.3

        cleared = client.patch(f"/api/v1/tanks/{tank_id}", json={"overflow_group": None}, headers=admin_auth_headers)
        assert cleared.json()["overflow_group"] is None
        assert cleared.json()["capacity_liters"] == 10000

    def test_regular_user_cannot_update(self, client, auth_headers, mazs):
        response = client.patch(f"/api/v1/tanks/{mazs['diesel'].id}", json={"capacity_liters": 1}, headers=auth_headers)
        assert response.status_code == 403

    def test_unknown_tank(self, client, admin_auth_headers, mazs):
        response = client.patch("/api/v1/tanks/99999", json={"capacity_liters": 1}, headers=admin_auth_headers)
        assert response.status_code == 404


class TestTankReadings:
    def test_readings_sorted_and_filtered(self, client, auth_headers, test_db, mazs):
        tank_id = mazs["diesel"].id
        base = datetime(2026, 9, 14, 10, 0)
        for i in range(5):
            test_db.add(TankReading(tank_id=tank_id, measured_at=base + timedelta(hours=i),
                                    volume=Decimal(5000 - i * 100), source_row_id=100 + i))
        test_db.commit()

        response = client.get(f"/api/v1/tanks/{tank_id}/readings",
                              params={"date_from": "2026-09-14T11:00:00", "limit": 2}, headers=auth_headers)
        assert response.status_code == 200
        body = response.json()
        assert body["total"] == 4
        assert [p["volume"] for p in body["items"]] == [4700.0, 4600.0]


class TestCardLimits:
    @pytest.fixture
    def limits(self, test_db, mazs):
        common = dict(provider_id=mazs["provider"].id, template_id=mazs["template"].id, synced_at=datetime(2026, 9, 14, 12, 0))
        rows = [
            CardLimit(source_card_id=1, source_fuel_id=3, card_code="A1", card_name="200 ИП Касумов 793", card_enabled=True,
                      fuel_type="АИ-92", limit_type_id=7, limit_type_name="Календарный день",
                      limit_liters=Decimal("40"), used_liters=Decimal("40"), **common),
            CardLimit(source_card_id=2, source_fuel_id=2, card_code="B2", card_name="УТ126", card_enabled=True,
                      fuel_type="ДТ", limit_type_id=7, limit_type_name="Календарный день",
                      limit_liters=Decimal("250"), used_liters=Decimal("225.17"), **common),
            CardLimit(source_card_id=3, source_fuel_id=2, card_code="C3", card_name="УТ110", card_enabled=True,
                      fuel_type="ДТ", limit_type_id=7, limit_type_name="Календарный день",
                      limit_liters=Decimal("150"), used_liters=Decimal("0"), **common),
            CardLimit(source_card_id=4, source_fuel_id=2, card_code="D4", card_name="К010", card_enabled=True,
                      fuel_type="ДТ", limit_type_id=0, limit_liters=Decimal("200"), used_liters=None, **common),
            CardLimit(source_card_id=5, source_fuel_id=3, card_code="E5", card_name="Выключена", card_enabled=False,
                      fuel_type="АИ-92", limit_type_id=7, limit_liters=Decimal("30"), used_liters=Decimal("30"), **common),
        ]
        test_db.add_all(rows)
        test_db.commit()
        return rows

    def test_sorted_by_usage_with_stats(self, client, auth_headers, limits):
        response = client.get("/api/v1/card-limits", headers=auth_headers)
        assert response.status_code == 200
        body = response.json()
        assert body["total"] == 4
        assert [item["card_name"] for item in body["items"][:2]] == ["200 ИП Касумов 793", "УТ126"]
        assert body["items"][1]["remaining_liters"] == 24.83
        assert body["items"][1]["used_percent"] == 90.1
        stats = body["stats"]
        assert stats == {"total": 5, "enabled": 4, "near_limit": 2, "exhausted": 1, "forbidden": 0, "without_period": 1}
        assert [p["name"] for p in body["providers"]] == ["МАЗС"]
        assert body["fuel_types"] == ["АИ-92", "ДТ"]

    def test_filters(self, client, auth_headers, limits):
        near = client.get("/api/v1/card-limits", params={"near_limit": True}, headers=auth_headers).json()
        assert {item["card_name"] for item in near["items"]} == {"200 ИП Касумов 793", "УТ126"}

        found = client.get("/api/v1/card-limits", params={"search": "УТ1", "fuel_type": "ДТ"}, headers=auth_headers).json()
        assert {item["card_name"] for item in found["items"]} == {"УТ126", "УТ110"}

        with_disabled = client.get("/api/v1/card-limits", params={"only_enabled": False}, headers=auth_headers).json()
        assert with_disabled["total"] == 5


class TestFillsByCard:
    def test_aggregates_days_and_compares_with_daily_limit(self, client, auth_headers, test_db, mazs):
        provider_id = mazs["provider"].id
        test_db.add(CardLimit(provider_id=provider_id, template_id=mazs["template"].id, source_card_id=1, source_fuel_id=3,
                              card_code="A1", card_name="214 ИП Касумов 772 ", card_enabled=True, fuel_type="АИ-92",
                              limit_type_id=7, limit_liters=Decimal("20")))
        fills = [
            ("214 ИП Касумов 772", "АИ-92", datetime(2026, 9, 10, 9, 0), "20", "1016201"),
            ("214 ИП Касумов 772", "АИ-92", datetime(2026, 9, 11, 9, 0), "10", "1016201"),
            ("214 ИП Касумов 772", "АИ-92", datetime(2026, 9, 11, 18, 0), "9", "1016201"),
            ("214 ИП Касумов 772", "АИ-92", datetime(2026, 9, 12, 9, 0), "12", "1016201"),
            ("УТ226", "ДТ", datetime(2026, 9, 11, 9, 0), "200", "807211"),
            ("УТ226", "ДТ", datetime(2026, 8, 1, 9, 0), "999", "807211"),  # вне периода
        ]
        for card, product, when, qty, azs in fills:
            test_db.add(Transaction(provider_id=provider_id, card_number=card, product=product, transaction_date=when,
                                    quantity=Decimal(qty), azs_number=azs))
        test_db.commit()

        other = Provider(name="РН-Карт", code="RN")
        test_db.add(other)
        test_db.flush()
        test_db.add(Transaction(provider_id=other.id, card_number="чужая карта", product="ДТ",
                                transaction_date=datetime(2026, 9, 11, 10, 0), quantity=Decimal("77"), azs_number="X1"))
        test_db.commit()

        response = client.get("/api/v1/reports/fills-by-card",
                              params={"date_from": "2026-09-10", "date_to": "2026-09-14"}, headers=auth_headers)
        assert response.status_code == 200
        body = response.json()
        # Провайдер без подключения к Топазу в отчёт блока «АЗС Топаз» не попадает
        assert body["totals"] == {"cards": 2, "fills_count": 5, "liters": 251.0}
        assert body["providers"] == [{"id": provider_id, "name": "МАЗС"}]
        assert body["fuel_types"] == ["АИ-92", "ДТ"]

        first = body["items"][0]
        assert first["card_number"] == "214 ИП Касумов 772"
        assert first["fills_count"] == 4
        assert first["days_with_fills"] == 3
        assert first["max_daily_liters"] == 20.0
        assert first["daily_limit"] == 20.0
        assert first["days_at_limit"] == 2
        assert first["azs_numbers"] == ["1016201"]

        diesel = body["items"][1]
        assert diesel["daily_limit"] is None
        assert diesel["days_at_limit"] is None

        at_limit = client.get("/api/v1/reports/fills-by-card",
                              params={"date_from": "2026-09-10", "date_to": "2026-09-14", "at_limit_only": True},
                              headers=auth_headers).json()
        assert [item["card_number"] for item in at_limit["items"]] == ["214 ИП Касумов 772"]

    def test_detail_lists_each_fill_of_card(self, client, auth_headers, test_db, mazs):
        provider_id = mazs["provider"].id
        other = Provider(name="РН-Карт", code="RN")
        test_db.add(other)
        test_db.flush()
        rows = [
            (provider_id, "214 ИП Касумов 772", "АИ-92", datetime(2026, 9, 11, 9, 0), "10", "1016201"),
            (provider_id, "214 ИП Касумов 772", "АИ-92", datetime(2026, 9, 11, 18, 0), "9", "1016201"),
            (provider_id, "УТ226", "ДТ", datetime(2026, 9, 12, 9, 0), "200", "807211"),
            (other.id, "214 ИП Касумов 772", "АИ-92", datetime(2026, 9, 12, 9, 0), "55", "X1"),  # не Топаз
        ]
        for prov, card, product, when, qty, azs in rows:
            test_db.add(Transaction(provider_id=prov, card_number=card, product=product, transaction_date=when,
                                    quantity=Decimal(qty), azs_number=azs))
        test_db.commit()

        body = client.get("/api/v1/reports/fills", params={
            "date_from": "2026-09-10", "date_to": "2026-09-14", "card_number": "214 ИП Касумов 772",
            "provider_id": provider_id, "fuel_type": "АИ-92",
        }, headers=auth_headers).json()
        assert body["total"] == 2
        assert body["truncated"] is False
        assert [(item["transaction_date"], item["liters"], item["azs_number"]) for item in body["items"]] == [
            ("2026-09-11T18:00:00", 9.0, "1016201"),
            ("2026-09-11T09:00:00", 10.0, "1016201"),
        ]

        everything = client.get("/api/v1/reports/fills", params={
            "date_from": "2026-09-10", "date_to": "2026-09-14", "limit": 2,
        }, headers=auth_headers).json()
        assert everything["total"] == 3
        assert everything["truncated"] is True
        assert everything["items"][0]["card_number"] == "УТ226"

    def test_xlsx_export_groups_fills_by_card(self, client, auth_headers, test_db, mazs):
        import io as _io
        from openpyxl import load_workbook

        provider_id = mazs["provider"].id
        test_db.add(CardLimit(provider_id=provider_id, template_id=mazs["template"].id, source_card_id=1, source_fuel_id=3,
                              card_code="A1", card_name="214 ИП Касумов 772", card_enabled=True, fuel_type="АИ-92",
                              limit_type_id=7, limit_liters=Decimal("20")))
        for card, product, when, qty in [
            ("214 ИП Касумов 772", "АИ-92", datetime(2026, 9, 11, 9, 0), "10"),
            ("214 ИП Касумов 772", "АИ-92", datetime(2026, 9, 11, 18, 0), "10"),
            ("214 ИП Касумов 772", "АИ-92", datetime(2026, 9, 12, 9, 0), "5"),
            ("УТ226", "ДТ", datetime(2026, 9, 12, 9, 0), "200"),
        ]:
            test_db.add(Transaction(provider_id=provider_id, card_number=card, product=product, transaction_date=when,
                                    quantity=Decimal(qty), azs_number="1016201", vehicle="УАЗ"))
        test_db.commit()

        response = client.get("/api/v1/reports/fills-by-card/export",
                              params={"date_from": "2026-09-10", "date_to": "2026-09-14"}, headers=auth_headers)
        assert response.status_code == 200
        assert "spreadsheetml" in response.headers["content-type"]
        assert "zapravki-karty_2026-09-10_2026-09-14.xlsx" in response.headers["content-disposition"]

        workbook = load_workbook(_io.BytesIO(response.content))
        assert workbook.sheetnames == ["Заправки по картам", "Итоги по картам"]
        summary = workbook["Итоги по картам"]
        assert [summary.cell(row=r, column=2).value for r in range(2, summary.max_row + 1)] == ["214 ИП Касумов 772", "УТ226"]

        details = workbook["Заправки по картам"]
        rows = list(details.iter_rows(min_row=2, values_only=True))
        assert rows[0][0].startswith("214 ИП Касумов 772 — 3 заправ., 2 дн.")
        assert [r[5] for r in rows[1:4]] == [10, 10, 5]
        assert [r[6] for r in rows[1:4]] == [20, 20, 5]
        assert [r[8] for r in rows[1:4]] == ["да", "да", None]
        assert details.row_dimensions[3].outlineLevel == 1
        assert details.row_dimensions[2].outlineLevel in (0, None)
        assert rows[4][0].startswith("УТ226")

        single = client.get("/api/v1/reports/fills-by-card/export", params={
            "date_from": "2026-09-10", "date_to": "2026-09-14", "card_number": "УТ226", "provider_id": provider_id,
        }, headers=auth_headers)
        single_book = load_workbook(_io.BytesIO(single.content))
        assert single_book["Итоги по картам"].max_row == 2
        assert "zapravki-karta_" in single.headers["content-disposition"]

    def test_rejects_long_or_inverted_period(self, client, auth_headers):
        assert client.get("/api/v1/reports/fills-by-card", params={"date_from": "2026-01-01", "date_to": "2026-09-01"},
                          headers=auth_headers).status_code == 400
        assert client.get("/api/v1/reports/fills-by-card", params={"date_from": "2026-09-10", "date_to": "2026-09-01"},
                          headers=auth_headers).status_code == 400


class TestCardLimitsGrouping:
    def test_same_limit_on_diesel_grades_is_one_row(self, client, auth_headers, test_db, mazs):
        common = dict(provider_id=mazs["provider"].id, template_id=mazs["template"].id, source_card_id=68,
                      card_code="0087D287", card_name="К068", card_enabled=True, fuel_type="ДТ",
                      limit_type_id=7, limit_type_name="Календарный день", limit_liters=Decimal("150"))
        test_db.add_all([
            CardLimit(source_fuel_id=1, source_fuel="ДТ1", used_liters=Decimal("0"), **common),
            CardLimit(source_fuel_id=5, source_fuel="ДТ2", used_liters=Decimal("0"), **common),
            CardLimit(source_fuel_id=6, source_fuel="ДТ3", used_liters=Decimal("120.05"), **common),
        ])
        test_db.commit()

        body = client.get("/api/v1/card-limits", params={"search": "к068"}, headers=auth_headers).json()
        assert body["total"] == 1
        item = body["items"][0]
        assert item["source_fuel"] == "ДТ1, ДТ2, ДТ3"
        assert item["used_liters"] == 120.05
        assert item["remaining_liters"] == 29.95
        assert body["stats"]["total"] == 1
