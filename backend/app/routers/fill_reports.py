"""
Отчёты по заправкам на АЗС Топаза на основе загруженных транзакций
"""
import io
from datetime import date, datetime, time, timedelta
from typing import Dict, List, Optional, Tuple

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.auth import require_auth_if_enabled
from app.database import get_db
from app.models import CardLimit, Provider, Transaction, User
from app.schemas import (
    FillDetailItem, FillsByCardItem, FillsByCardResponse, FillsByCardTotals, FillsDetailResponse, TopazProviderOption,
)
from app.services.topaz_sync_service import LIMIT_TYPE_DAY, topaz_providers

router = APIRouter(prefix="/api/v1/reports", tags=["reports"])

MAX_PERIOD_DAYS = 93
AT_LIMIT_SHARE = 0.9
EXPORT_MAX_FILLS = 50000

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
LITERS_FORMAT = "#,##0.00"
DATETIME_FORMAT = "DD.MM.YYYY HH:MM"


def _card_key(value: Optional[str]) -> str:
    return (value or "").strip().lower()


def _period(date_from: Optional[date], date_to: Optional[date]) -> Tuple[date, date]:
    date_to = date_to or date.today()
    date_from = date_from or (date_to - timedelta(days=29))
    if date_from > date_to:
        raise HTTPException(status_code=400, detail="Дата «с» позже даты «по»")
    if (date_to - date_from).days + 1 > MAX_PERIOD_DAYS:
        raise HTTPException(status_code=400, detail=f"Период не больше {MAX_PERIOD_DAYS} дней")
    return date_from, date_to


def _scope_filters(date_from: date, date_to: date, scope_ids: List[int], card_number: Optional[str] = None,
                   azs_number: Optional[str] = None, fuel_type: Optional[str] = None, search: Optional[str] = None) -> list:
    """Условия на заправки: период, провайдеры Топаза и фильтры отчёта."""
    conditions = [
        Transaction.transaction_date >= datetime.combine(date_from, time.min),
        Transaction.transaction_date <= datetime.combine(date_to, time.max),
        Transaction.quantity > 0,
        Transaction.provider_id.in_(scope_ids),
    ]
    if card_number is not None:
        conditions.append(Transaction.card_number == card_number)
    if azs_number:
        conditions.append(Transaction.azs_number == azs_number)
    if fuel_type:
        conditions.append(Transaction.product == fuel_type)
    if search:
        conditions.append(Transaction.card_number.ilike(f"%{search.strip()}%"))
    return conditions


def _daily_limits(db: Session) -> Dict[Tuple, float]:
    limits: Dict[Tuple, float] = {}
    for prov, card_name, fuel, limit_liters in db.query(
        CardLimit.provider_id, CardLimit.card_name, CardLimit.fuel_type, CardLimit.limit_liters
    ).filter(CardLimit.limit_type_id == LIMIT_TYPE_DAY, CardLimit.limit_liters > 0):
        limits[(prov, _card_key(card_name), fuel)] = float(limit_liters)
    return limits


def _summary_items(db: Session, conditions: list, at_limit_only: bool, daily_limits: Dict[Tuple, float],
                   providers: Dict[int, str]) -> List[FillsByCardItem]:
    """Итоги по картам: литры, дни с заправками, максимум в сутки и дни у суточного лимита."""
    day = func.date(Transaction.transaction_date)
    daily_rows = db.query(
        Transaction.provider_id,
        Transaction.card_number,
        Transaction.product,
        day.label("day"),
        func.count(Transaction.id).label("fills"),
        func.sum(Transaction.quantity).label("liters"),
        func.min(Transaction.transaction_date).label("first_fill"),
        func.max(Transaction.transaction_date).label("last_fill"),
    ).filter(*conditions).group_by(Transaction.provider_id, Transaction.card_number, Transaction.product, day).all()

    azs_by_key: Dict[Tuple, set] = {}
    for prov, card, product, azs in db.query(
        Transaction.provider_id, Transaction.card_number, Transaction.product, Transaction.azs_number
    ).filter(*conditions).distinct():
        if azs:
            azs_by_key.setdefault((prov, card, product), set()).add(azs)

    aggregated: Dict[Tuple, dict] = {}
    for row in daily_rows:
        key = (row.provider_id, row.card_number, row.product)
        entry = aggregated.setdefault(key, {
            "fills": 0, "liters": 0.0, "days": 0, "max_daily": 0.0,
            "first": None, "last": None, "daily": [],
        })
        liters = float(row.liters or 0)
        entry["fills"] += int(row.fills or 0)
        entry["liters"] += liters
        entry["days"] += 1
        entry["max_daily"] = max(entry["max_daily"], liters)
        entry["daily"].append(liters)
        entry["first"] = row.first_fill if entry["first"] is None else min(entry["first"], row.first_fill)
        entry["last"] = row.last_fill if entry["last"] is None else max(entry["last"], row.last_fill)

    items = []
    for (prov, card, product), entry in aggregated.items():
        daily_limit = daily_limits.get((prov, _card_key(card), product))
        days_at_limit = None
        if daily_limit:
            days_at_limit = sum(1 for liters in entry["daily"] if liters >= daily_limit * AT_LIMIT_SHARE)
        if at_limit_only and not days_at_limit:
            continue
        items.append(FillsByCardItem(
            provider_id=prov,
            provider_name=providers.get(prov),
            card_number=card,
            fuel_type=product,
            fills_count=entry["fills"],
            liters=round(entry["liters"], 2),
            days_with_fills=entry["days"],
            max_daily_liters=round(entry["max_daily"], 2),
            avg_daily_liters=round(entry["liters"] / entry["days"], 2) if entry["days"] else 0.0,
            first_fill=entry["first"],
            last_fill=entry["last"],
            azs_numbers=sorted(azs_by_key.get((prov, card, product), set())),
            daily_limit=daily_limit,
            days_at_limit=days_at_limit,
        ))
    items.sort(key=lambda item: (-(item.days_at_limit or 0), -item.liters))
    return items


def _provider_names(db: Session) -> Dict[int, str]:
    return {pid: name for pid, name in db.query(Provider.id, Provider.name)}


@router.get("/fills", response_model=FillsDetailResponse)
def fills_detail(
    date_from: Optional[date] = Query(None, description="С (по умолчанию — 30 дней назад)"),
    date_to: Optional[date] = Query(None, description="По (по умолчанию — сегодня)"),
    provider_id: Optional[int] = Query(None),
    card_number: Optional[str] = Query(None, description="Точное название карты — детализация одной строки отчёта"),
    azs_number: Optional[str] = Query(None),
    fuel_type: Optional[str] = Query(None),
    search: Optional[str] = Query(None, description="Поиск по карте"),
    limit: int = Query(5000, ge=1, le=20000),
    db: Session = Depends(get_db),
    _: Optional[User] = Depends(require_auth_if_enabled),
):
    """Детализация «Заправок по картам»: каждая заправка на АЗС Топаза, самые свежие первыми."""
    date_from, date_to = _period(date_from, date_to)
    scope_ids = [pid for pid, _name in topaz_providers(db) if not provider_id or pid == provider_id]
    query = db.query(Transaction).filter(
        *_scope_filters(date_from, date_to, scope_ids, card_number, azs_number, fuel_type, search)
    )
    total = query.count()
    rows = query.order_by(Transaction.transaction_date.desc(), Transaction.id.desc()).limit(limit).all()
    providers = _provider_names(db)
    return FillsDetailResponse(
        date_from=date_from,
        date_to=date_to,
        total=total,
        truncated=total > len(rows),
        items=[
            FillDetailItem(
                id=row.id,
                transaction_date=row.transaction_date,
                provider_id=row.provider_id,
                provider_name=providers.get(row.provider_id),
                card_number=row.card_number,
                vehicle=row.vehicle,
                azs_number=row.azs_number,
                fuel_type=row.product,
                liters=float(row.quantity),
            )
            for row in rows
        ],
    )


@router.get("/fills-by-card", response_model=FillsByCardResponse)
def fills_by_card(
    date_from: Optional[date] = Query(None, description="С (по умолчанию — 30 дней назад)"),
    date_to: Optional[date] = Query(None, description="По (по умолчанию — сегодня)"),
    provider_id: Optional[int] = Query(None),
    azs_number: Optional[str] = Query(None),
    fuel_type: Optional[str] = Query(None),
    search: Optional[str] = Query(None, description="Поиск по карте"),
    at_limit_only: bool = Query(False, description="Только карты, которые упирались в суточный лимит"),
    db: Session = Depends(get_db),
    _: Optional[User] = Depends(require_auth_if_enabled),
):
    """Заправки по картам за период: литры, дни с заправками, максимум в сутки и сравнение с суточным лимитом."""
    date_from, date_to = _period(date_from, date_to)
    # Отчёт — часть блока «АЗС Топаз»: только провайдеры, чьи АЗС читаются из Топаза
    topaz = topaz_providers(db)
    topaz_ids = [pid for pid, _name in topaz]
    scope_ids = [pid for pid in topaz_ids if not provider_id or pid == provider_id]

    items = _summary_items(
        db, _scope_filters(date_from, date_to, scope_ids, None, azs_number, fuel_type, search),
        at_limit_only, _daily_limits(db), _provider_names(db),
    )
    return FillsByCardResponse(
        date_from=date_from,
        date_to=date_to,
        total=len(items),
        items=items,
        totals=FillsByCardTotals(
            cards=len({(item.provider_id, item.card_number) for item in items}),
            fills_count=sum(item.fills_count for item in items),
            liters=round(sum(item.liters for item in items), 2),
        ),
        providers=[TopazProviderOption(id=pid, name=name) for pid, name in topaz],
        fuel_types=[
            product for (product,) in db.query(Transaction.product)
            .filter(Transaction.provider_id.in_(topaz_ids), Transaction.product.isnot(None))
            .distinct()
            .order_by(Transaction.product)
        ] if topaz_ids else [],
    )


# --------------------------------------------------------------------------- XLSX

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
_GROUP_FONT = Font(bold=True)
_GROUP_FILL = PatternFill(start_color="EEF1F6", end_color="EEF1F6", fill_type="solid")
_WARN_FONT = Font(bold=True, color="A35A06")


def _write_header(ws, headers: List[str], widths: List[int]) -> None:
    ws.append(headers)
    for index, width in enumerate(widths, start=1):
        cell = ws.cell(row=1, column=index)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"


def _summary_sheet(ws, items: List[FillsByCardItem]) -> None:
    _write_header(ws, [
        "Провайдер", "Карта", "Топливо", "Заправок", "Литров", "Дней с заправками", "Максимум за сутки, л",
        "Среднее за сутки, л", "Суточный лимит, л", "Дней у лимита", "АЗС", "Первая заправка", "Последняя заправка",
    ], [12, 30, 10, 11, 13, 12, 13, 13, 12, 11, 18, 17, 17])
    for item in items:
        ws.append([
            item.provider_name, item.card_number, item.fuel_type, item.fills_count, item.liters, item.days_with_fills,
            item.max_daily_liters, item.avg_daily_liters, item.daily_limit, item.days_at_limit,
            ", ".join(item.azs_numbers), item.first_fill, item.last_fill,
        ])
        row = ws.max_row
        for column in (5, 7, 8, 9):
            ws.cell(row=row, column=column).number_format = LITERS_FORMAT
        for column in (12, 13):
            ws.cell(row=row, column=column).number_format = DATETIME_FORMAT
        if item.days_at_limit:
            ws.cell(row=row, column=10).font = _WARN_FONT
    ws.auto_filter.ref = f"A1:M{max(ws.max_row, 1)}"


def _details_sheet(ws, items: List[FillsByCardItem], fills: List[Transaction]) -> None:
    """Заправки, сгруппированные по карте: строка итога и под ней свёртываемые строки заправок."""
    ws.sheet_properties.outlinePr.summaryBelow = False
    _write_header(ws, [
        "Карта / дата и время", "Провайдер", "АЗС", "Закреплена за", "Топливо", "Литров", "За день, л", "Суточный лимит, л",
        "Лимит выбран",
    ], [30, 12, 12, 26, 10, 12, 12, 12, 13])

    by_card: Dict[Tuple, List[Transaction]] = {}
    for fill in fills:
        by_card.setdefault((fill.provider_id, fill.card_number, fill.product), []).append(fill)

    for item in items:
        card_fills = sorted(by_card.get((item.provider_id, item.card_number, item.fuel_type), []),
                            key=lambda f: (f.transaction_date, f.id))
        ws.append([
            f"{item.card_number or 'без карты'} — {item.fills_count} заправ., {item.days_with_fills} дн.",
            item.provider_name, ", ".join(item.azs_numbers), None, item.fuel_type, item.liters, None, item.daily_limit,
            f"{item.days_at_limit} дн." if item.days_at_limit else None,
        ])
        group_row = ws.max_row
        for column in range(1, 10):
            cell = ws.cell(row=group_row, column=column)
            cell.font = _GROUP_FONT
            cell.fill = _GROUP_FILL
        ws.cell(row=group_row, column=6).number_format = LITERS_FORMAT
        ws.cell(row=group_row, column=8).number_format = LITERS_FORMAT

        day_totals: Dict[date, float] = {}
        for fill in card_fills:
            day_totals[fill.transaction_date.date()] = day_totals.get(fill.transaction_date.date(), 0.0) + float(fill.quantity)
        for fill in card_fills:
            day_total = round(day_totals[fill.transaction_date.date()], 2)
            at_limit = bool(item.daily_limit) and day_total >= item.daily_limit * AT_LIMIT_SHARE
            ws.append([
                fill.transaction_date, None, fill.azs_number, fill.vehicle, fill.product, float(fill.quantity),
                day_total, item.daily_limit, "да" if at_limit else None,
            ])
            row = ws.max_row
            ws.row_dimensions[row].outlineLevel = 1
            ws.cell(row=row, column=1).number_format = DATETIME_FORMAT
            ws.cell(row=row, column=1).alignment = Alignment(indent=1, horizontal="left")
            for column in (6, 7, 8):
                ws.cell(row=row, column=column).number_format = LITERS_FORMAT
            if at_limit:
                ws.cell(row=row, column=9).font = _WARN_FONT


@router.get("/fills-by-card/export")
def export_fills_by_card(
    date_from: Optional[date] = Query(None, description="С (по умолчанию — 30 дней назад)"),
    date_to: Optional[date] = Query(None, description="По (по умолчанию — сегодня)"),
    provider_id: Optional[int] = Query(None),
    card_number: Optional[str] = Query(None, description="Только эта карта — выгрузка из детализации"),
    azs_number: Optional[str] = Query(None),
    fuel_type: Optional[str] = Query(None),
    search: Optional[str] = Query(None, description="Поиск по карте"),
    at_limit_only: bool = Query(False, description="Только карты, которые упирались в суточный лимит"),
    db: Session = Depends(get_db),
    _: Optional[User] = Depends(require_auth_if_enabled),
):
    """XLSX: лист итогов по картам и лист заправок, сгруппированных по карте."""
    date_from, date_to = _period(date_from, date_to)
    scope_ids = [pid for pid, _name in topaz_providers(db) if not provider_id or pid == provider_id]
    conditions = _scope_filters(date_from, date_to, scope_ids, card_number, azs_number, fuel_type, search)

    items = _summary_items(db, conditions, at_limit_only, _daily_limits(db), _provider_names(db))
    fills_query = db.query(Transaction).filter(*conditions)
    if fills_query.count() > EXPORT_MAX_FILLS:
        raise HTTPException(status_code=400, detail=f"Больше {EXPORT_MAX_FILLS} заправок — сузьте период или фильтры")
    fills = fills_query.all()

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Итоги по картам"
    _summary_sheet(summary, items)
    _details_sheet(workbook.create_sheet("Заправки по картам"), items, fills)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    suffix = "karta" if card_number is not None else "karty"
    filename = f"zapravki-{suffix}_{date_from.isoformat()}_{date_to.isoformat()}.xlsx"
    return StreamingResponse(output, media_type=XLSX_MEDIA_TYPE, headers={"Content-Disposition": f"attachment; filename={filename}"})
